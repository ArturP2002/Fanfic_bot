"""Импорт стартовой базы из CSV / Excel."""

from __future__ import annotations

import asyncio
import csv
import logging
from pathlib import Path

import httpx
from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Work, WorkAlias
from app.models.enums import ConfirmationStatus
from app.services.normalize import normalize_text, split_aliases
from app.services.parsers import ParsedWork, canonical_url, detect_platform, parse_fic_url
from app.services.parsers.base import WorkPageUnavailableError
from app.services.works import (
    create_work,
    enrich_work,
    find_work_by_url,
    parse_work_status,
    split_translation_urls,
    volume_from_words,
)

logger = logging.getLogger(__name__)

# Пауза между загрузками страниц, чтобы не нагружать площадки и не ловить бан
FETCH_PAUSE_SECONDS = 1.5

# Заголовки шаблона (русские и английские варианты)
FIELD_MAP = {
    "название": "title",
    "title": "title",
    "автор": "author",
    "author": "author",
    "оригинальное название": "original_title",
    "original_title": "original_title",
    "русское название": "russian_title",
    "russian_title": "russian_title",
    "другие названия": "other_titles",
    "other_titles": "other_titles",
    "сокращения": "abbreviations",
    "abbreviations": "abbreviations",
    "фандом": "fandom",
    "fandom": "fandom",
    "пейринг": "pairing",
    "pairing": "pairing",
    "статус": "status",
    "status": "status",
    "объём": "volume",
    "объем": "volume",
    "volume": "volume",
    "words": "words",
    "слов": "words",
    "ссылка": "source_url",
    "source_url": "source_url",
    "ссылка на перевод": "translation_url",
    "translation_url": "translation_url",
    "теги": "tags",
    "tags": "tags",
}


def _normalize_header(header: str) -> str:
    return header.strip().lower()


def _map_row(headers: list[str], values: list[str]) -> dict[str, str]:
    mapped: dict[str, str] = {}
    for header, value in zip(headers, values, strict=False):
        key = FIELD_MAP.get(_normalize_header(header))
        if key and value is not None:
            mapped[key] = str(value).strip()
    return mapped


def _parse_tags(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [part.strip() for part in raw.replace(";", ",").split(",") if part.strip()]


def _parse_words(raw: str | None) -> int | None:
    if not raw:
        return None
    digits = "".join(ch for ch in raw if ch.isdigit())
    return int(digits) if digits else None


def read_table(path: Path) -> list[dict[str, str]]:
    """Читает CSV или XLSX в список словарей с каноническими ключами."""
    suffix = path.suffix.lower()
    rows: list[dict[str, str]] = []

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            headers = next(reader, None)
            if not headers:
                return []
            for raw in reader:
                if not any(cell.strip() for cell in raw):
                    continue
                rows.append(_map_row(headers, raw))
    elif suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        headers_raw = next(iterator, None)
        if not headers_raw:
            return []
        headers = [str(h or "") for h in headers_raw]
        for raw in iterator:
            values = ["" if v is None else str(v) for v in raw]
            if not any(v.strip() for v in values):
                continue
            rows.append(_map_row(headers, values))
    else:
        raise ValueError(f"Неподдерживаемый формат файла: {suffix}")

    return rows


def _fill_from_parsed(
    row: dict[str, str],
    parsed: ParsedWork,
    original: ParsedWork | None,
) -> None:
    """Дополняет строку таблицы метаданными со страницы; заполненное вручную не трогаем.

    У переводов автор и оригинальное название берутся со страницы оригинала.
    """
    from_page = {
        "title": parsed.title,
        "author": (original.author if original else None) or parsed.author,
        "original_title": original.title if original else None,
        "fandom": parsed.fandom,
        "pairing": parsed.pairing,
        "status": parsed.status,
        "words": str(parsed.word_count) if parsed.word_count else None,
    }
    for key, value in from_page.items():
        if value and not row.get(key):
            row[key] = value


async def _find_existing(
    session: AsyncSession,
    row: dict[str, str],
    urls: list[str | None],
) -> Work | None:
    """Ищет уже существующую работу: по ссылкам, затем по названиям с учётом автора."""
    for url in urls:
        if url:
            work = await find_work_by_url(session, url)
            if work:
                return work

    norm_author = normalize_text(row.get("author"))
    for name in (row.get("title"), row.get("original_title"), row.get("russian_title")):
        norm = normalize_text(name)
        if not norm:
            continue
        result = await session.execute(
            select(Work)
            .outerjoin(WorkAlias, WorkAlias.work_id == Work.id)
            .where(or_(Work.normalized_title == norm, WorkAlias.normalized_alias == norm))
        )
        for work in result.unique().scalars().all():
            if not norm_author or not work.normalized_author or work.normalized_author == norm_author:
                return work
    return None


async def import_works_from_file(
    session: AsyncSession,
    path: Path,
) -> dict[str, int]:
    """Импортирует произведения. Повторный импорт не создаёт дубли.

    Если в строке есть только ссылка на AO3 / Ficbook / FanFiction.net,
    название, автор и остальные поля подтягиваются со страницы. У переводов
    дополнительно загружается оригинал, чтобы перевод и оригинал были одной работой.
    """
    rows = read_table(path)
    stats = {"total": len(rows), "created": 0, "updated": 0, "skipped": 0, "failed": 0}
    fetched_any = False

    async def fetch(url: str) -> ParsedWork | None:
        nonlocal fetched_any
        if fetched_any:
            await asyncio.sleep(FETCH_PAUSE_SECONDS)
        fetched_any = True
        try:
            return await parse_fic_url(url)
        except (httpx.HTTPError, WorkPageUnavailableError) as exc:
            logger.warning("Страница недоступна при импорте: %s (%s)", url, exc)
            return None
        except Exception:
            logger.exception("Не удалось загрузить страницу при импорте: %s", url)
            return None

    for row in rows:
        title = row.get("title", "").strip()
        source_url = row.get("source_url") or None
        if source_url and detect_platform(source_url):
            source_url = canonical_url(source_url)
        translation_url = row.get("translation_url") or None
        parsed: ParsedWork | None = None

        if not title and source_url and detect_platform(source_url):
            if await find_work_by_url(session, source_url):
                stats["skipped"] += 1
                logger.info("Пропуск дубля при импорте по ссылке: %s", source_url)
                continue
            parsed = await fetch(source_url)
            if not parsed:
                stats["failed"] += 1
                continue
            original = None
            if parsed.original_url and detect_platform(parsed.original_url):
                original = await fetch(canonical_url(parsed.original_url))
            _fill_from_parsed(row, parsed, original)
            title = row.get("title", "").strip()
            source_url, page_translation_url = split_translation_urls(
                source_url, parsed.original_url
            )
            translation_url = translation_url or page_translation_url

        if not title:
            stats["skipped"] += 1
            continue

        author = row.get("author") or None
        existing = await _find_existing(session, row, [source_url, translation_url])
        if existing:
            extra_titles = [title, *split_aliases(row.get("other_titles"))]
            extra_titles += split_aliases(row.get("abbreviations"))
            enriched = parsed is not None and await enrich_work(
                session,
                existing,
                fandom=row.get("fandom") or None,
                pairing=row.get("pairing") or None,
                status=parse_work_status(row.get("status")),
                word_count=_parse_words(row.get("words")),
                tags=parsed.tags,
                source_url=source_url,
                translation_url=translation_url,
                original_title=row.get("original_title") or None,
                russian_title=row.get("russian_title") or None,
                extra_titles=extra_titles,
            )
            stats["updated" if enriched else "skipped"] += 1
            logger.info(
                "%s при импорте: «%s» → work_id=%s",
                "Дополнена существующая работа" if enriched else "Пропуск дубля",
                title,
                existing.id,
            )
            continue

        words = _parse_words(row.get("words") or row.get("volume"))
        volume_label = row.get("volume") if row.get("volume") and not words else volume_from_words(words)

        await create_work(
            session,
            title=title,
            author=author,
            original_title=row.get("original_title") or None,
            russian_title=row.get("russian_title") or None,
            fandom=row.get("fandom") or None,
            pairing=row.get("pairing") or None,
            status=parse_work_status(row.get("status")),
            word_count=words,
            volume_label=volume_label,
            tags=_parse_tags(row.get("tags")) or (parsed.tags if parsed else []),
            source_url=source_url,
            translation_url=translation_url,
            confirmation_status=ConfirmationStatus.CONFIRMED.value,
            other_titles=row.get("other_titles"),
            abbreviations=row.get("abbreviations"),
        )
        stats["created"] += 1

    logger.info(
        "Импорт завершён: всего=%s создано=%s дополнено=%s пропущено=%s ошибок=%s",
        stats["total"],
        stats["created"],
        stats["updated"],
        stats["skipped"],
        stats["failed"],
    )
    return stats
